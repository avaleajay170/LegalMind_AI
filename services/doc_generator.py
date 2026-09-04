import os
import uuid
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.gemini_service import ask_gemini_once
from config import Config


# ── DOCUMENT PROMPTS ──────────────────────────────────────
DOC_PROMPTS = {

    "Legal Notice": """Draft a formal Legal Notice for Indian courts.

Case Details:
- Client Name     : {client_name}
- Opposing Party  : {opposing_party}
- Case Type       : {case_type}
- Court           : {court_name}
- Description     : {description}
- Advocate        : {lawyer_name} ({enrollment_no})
- Date            : {date}

Write a complete, formal Legal Notice with:
1. Sender and recipient details
2. Subject line
3. Facts of the matter
4. Legal basis and sections
5. Relief sought
6. Consequences of non-compliance
7. Time limit for response (usually 15-30 days)
8. Advocate signature block

Use formal Indian legal language. Return only the document text.""",


    "FIR Draft": """Draft a First Information Report (FIR) for Indian police.

Case Details:
- Complainant     : {client_name}
- Case Type       : {case_type}
- Court/Station   : {court_name}
- Description     : {description}
- Advocate        : {lawyer_name} ({enrollment_no})
- Date            : {date}

Write a complete FIR draft with:
1. Police station details
2. Complainant information
3. Date, time and place of incident
4. Detailed facts of the incident
5. Names of accused (if known)
6. Witnesses (if any)
7. Applicable IPC sections
8. Relief requested
9. Declaration by complainant

Use formal language suitable for Indian police records.""",


    "Affidavit": """Draft a formal Affidavit for Indian courts.

Case Details:
- Deponent Name   : {client_name}
- Case Type       : {case_type}
- Court           : {court_name}
- Description     : {description}
- Advocate        : {lawyer_name} ({enrollment_no})
- Date            : {date}

Write a complete Affidavit with:
1. Court header
2. Deponent details and address
3. Verification statement
4. Numbered paragraphs of facts
5. Exhibits list if applicable
6. Deponent signature block
7. Notary/Commissioner section

Use standard Indian court affidavit format.""",


    "Bail Application": """Draft a Bail Application for Indian courts.

Case Details:
- Accused Name    : {client_name}
- Case Type       : {case_type}
- Court           : {court_name}
- Description     : {description}
- Advocate        : {lawyer_name} ({enrollment_no})
- Date            : {date}

Write a complete Bail Application with:
1. Court address and case details
2. Applicant/accused details
3. Nature of offence and sections
4. Grounds for bail (at least 5 strong grounds)
5. Personal details and ties to community
6. Surety details
7. Conditions the accused is willing to comply with
8. Prayer/relief sought
9. Advocate signature

Make grounds compelling and legally sound.""",


    "Contract Agreement": """Draft a formal Contract Agreement under Indian law.

Case Details:
- Party 1 (Client) : {client_name}
- Case Type        : {case_type}
- Court/Jurisdiction: {court_name}
- Description      : {description}
- Advocate         : {lawyer_name} ({enrollment_no})
- Date             : {date}

Write a complete Contract Agreement with:
1. Parties to the contract
2. Recitals/Background
3. Definitions
4. Scope of agreement
5. Terms and conditions (at least 8 clauses)
6. Payment terms if applicable
7. Duration and termination
8. Dispute resolution (arbitration clause)
9. Governing law (Indian law)
10. Signature blocks for all parties

Use formal legal language compliant with Indian Contract Act 1872."""
}


# ── GENERATE DOCUMENT CONTENT ─────────────────────────────
def generate_doc_content(template_type, case, lawyer_name, enrollment_no):
    """
    Generate document content using Gemini.

    Args:
        template_type : type of document to generate
        case          : case dict from MongoDB
        lawyer_name   : advocate's full name
        enrollment_no : Bar Council enrollment number

    Returns:
        str — generated document text
    """
    prompt_template = DOC_PROMPTS.get(template_type)

    if not prompt_template:
        return None

    prompt = prompt_template.format(
        client_name   = case.get("client_name",  "Client Name"),
        opposing_party= case.get("opposing_party","Opposing Party"),
        case_type     = case.get("case_type",     "General"),
        court_name    = case.get("court_name",    "Hon'ble Court"),
        description   = case.get("description",   ""),
        lawyer_name   = lawyer_name,
        enrollment_no = enrollment_no,
        date          = datetime.now().strftime("%d %B %Y")
    )

    return ask_gemini_once(prompt)


# ── CREATE DOCX FILE ──────────────────────────────────────
def create_docx(content, template_type, case_title):
    """
    Create a DOCX file from generated content.

    Args:
        content       : generated document text
        template_type : type of document
        case_title    : case title for filename

    Returns:
        dict with filepath and filename
    """
    doc = Document()

    # ── PAGE SETUP ────────────────────────────────────────
    section = doc.sections[0]
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)

    # ── HEADER ────────────────────────────────────────────
    header = doc.add_heading("LegalMind AI — Generated Document", level=1)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Document type
    doc_type_para = doc.add_paragraph(template_type.upper())
    doc_type_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_type_run  = doc_type_para.runs[0]
    doc_type_run.bold     = True
    doc_type_run.font.size= Pt(14)

    # Date
    date_para = doc.add_paragraph(
        f"Generated on: {datetime.now().strftime('%d %B %Y')}"
    )
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")  # spacer

    # ── CONTENT ───────────────────────────────────────────
    lines = content.split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph("")
            continue

        # Detect headings
        if line.isupper() and len(line) < 60:
            p = doc.add_heading(line, level=2)
        elif line.startswith(("1.", "2.", "3.", "4.", "5.",
                               "6.", "7.", "8.", "9.", "10.")):
            p = doc.add_paragraph(line, style="List Number")
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    # ── FOOTER ────────────────────────────────────────────
    doc.add_paragraph("")
    footer = doc.add_paragraph(
        "This document was generated by LegalMind AI. "
        "Please review and verify before use."
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size   = Pt(9)
    footer.runs[0].font.color.rgb = None

    # ── SAVE FILE ─────────────────────────────────────────
    os.makedirs(Config.GENERATED_FOLDER, exist_ok=True)

    safe_title = case_title.replace(" ", "_")[:30]
    safe_type  = template_type.replace(" ", "_")
    unique_id  = str(uuid.uuid4())[:6]
    filename   = f"{safe_type}_{safe_title}_{unique_id}.docx"
    filepath   = os.path.join(Config.GENERATED_FOLDER, filename)

    doc.save(filepath)
    return {
        "filename": filename,
        "filepath": filepath
    }


# ── MAIN GENERATE FUNCTION ────────────────────────────────
def generate_document(template_type, case, lawyer_name, enrollment_no):
    """
    Generate a complete legal document.

    Args:
        template_type : Legal Notice | FIR Draft | Affidavit |
                        Bail Application | Contract Agreement
        case          : case dict from MongoDB
        lawyer_name   : advocate's full name
        enrollment_no : Bar Council enrollment number

    Returns:
        dict with keys:
            success  (bool)
            filename (str)
            filepath (str)
            content  (str)
            message  (str)
    """
    if template_type not in Config.TEMPLATES:
        return {
            "success": False,
            "message": f"Invalid template type: {template_type}"
        }

    # Generate content using Gemini
    print(f"⏳ Generating {template_type}...")
    content = generate_doc_content(
        template_type = template_type,
        case          = case,
        lawyer_name   = lawyer_name,
        enrollment_no = enrollment_no
    )

    if not content:
        return {
            "success": False,
            "message": "AI could not generate document. Please try again."
        }

    # Create DOCX file
    file_info = create_docx(
        content       = content,
        template_type = template_type,
        case_title    = case.get("title", "Case")
    )

    print(f"✅ Document generated: {file_info['filename']}")

    return {
        "success":  True,
        "filename": file_info["filename"],
        "filepath": file_info["filepath"],
        "content":  content,
        "message":  f"{template_type} generated successfully!"
    } 
